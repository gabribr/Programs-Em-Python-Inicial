from docx import Document
import zipfile
from io import BytesIO
from datetime import datetime
import re

def sanitize_filename(filename):
    return re.sub(r'[<>:"/\\|?*]', '_', filename)

schools = [f"Escola Número {i+1}" for i in range(200)]

with zipfile.ZipFile('escolas.zip', 'w', zipfile.ZIP_DEFLATED) as zf:
    for i, school in enumerate(schools):
        doc = Document()
        doc.add_heading(school, 0)
        
        info_para = doc.add_paragraph(
            f"Informações da escola: {school}. "
            f"Data: {datetime.now().strftime('%d/%m/%Y')}. "
            f"Número sequencial: {i+1}."
        )
        
        bio = BytesIO()
        doc.save(bio)
        bio.seek(0)
        
        sanitized_name = sanitize_filename(f"{school}.docx")
        zf.writestr(sanitized_name, bio.read())

print("ZIP 'escolas.zip' criado com sucesso com 200 arquivos Word!")